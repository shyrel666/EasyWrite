import re
from collections import Counter
from typing import List, Dict, Any, Optional
from pathlib import Path
import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

class ParsedSection:
    def __init__(self, section_id: str, title: str, level: int, breadcrumb: str):
        self.section_id = section_id
        self.title = title.strip()
        self.level = level
        self.breadcrumb = breadcrumb
        self.paragraphs: List[str] = []
        self.tables: List[List[List[str]]] = []  # table rows and cells
        self.subsections: List['ParsedSection'] = []

    def add_paragraph(self, text: str):
        cleaned = text.strip()
        if cleaned:
            self.paragraphs.append(cleaned)

    def add_table(self, table_data: List[List[str]]):
        if table_data:
            self.tables.append(table_data)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "section_id": self.section_id,
            "title": self.title,
            "level": self.level,
            "breadcrumb": self.breadcrumb,
            "content": "\n\n".join(self.paragraphs),
            "tables": self.tables,
            "subsections": [sub.to_dict() for sub in self.subsections]
        }


class DocumentTreeBuilder:
    """
    按文档顺序接收「标题 / 正文段落 / 表格」，构建章节树、按序全文与扁平表格列表。
    Word 与 PDF 解析器共用，保证下游（拆标、评分表抽取、偏离表、知识库分块）拿到同构的结果。
    """

    def __init__(self):
        # 默认起始段落如果没有任何标题，放入前言/概述
        self.current = ParsedSection(section_id="sec_0", title="前言/基本信息", level=1, breadcrumb="前言/基本信息")
        self.roots: List[ParsedSection] = [self.current]
        # 栈结构，保存当前层级的父节点: [(level, section_obj)]
        self.stack: List[tuple] = [(1, self.current)]
        self.counter = 1
        self.text_lines: List[str] = []
        self.tables: List[Dict[str, Any]] = []
        self.recent_paragraphs: List[str] = []

    def add_heading(self, text: str, level: int):
        # 弹出比当前级别高或相等的栈顶
        while self.stack and self.stack[-1][0] >= level:
            self.stack.pop()
        breadcrumb = " > ".join([s.title for _, s in self.stack] + [text])
        section = ParsedSection(section_id=f"sec_{self.counter}", title=text, level=level, breadcrumb=breadcrumb)
        self.counter += 1
        if self.stack:
            # 挂载为父级子节点
            self.stack[-1][1].subsections.append(section)
        else:
            self.roots.append(section)
        self.stack.append((level, section))
        self.current = section
        self.recent_paragraphs = []
        self.text_lines.append(text)

    def add_paragraph(self, text: str):
        self.current.add_paragraph(text)
        self.recent_paragraphs = (self.recent_paragraphs + [text])[-3:]
        self.text_lines.append(text)

    def add_table(self, rows: List[List[str]], cell_keys: Optional[List[List[int]]] = None):
        """rows 为单元格文本；cell_keys 为同形状的物理单元格编号（纵向合并单元格各行编号相同）"""
        if not rows:
            return
        self.current.add_table(rows)
        self.tables.append({
            "index": len(self.tables),
            "breadcrumb": self.current.breadcrumb,
            "context_before": "\n".join(self.recent_paragraphs),
            "rows": rows,
            "cell_keys": cell_keys or [],
        })
        self.text_lines.extend(" | ".join(c for c in row if c) for row in rows)

    def result(self, doc_path: str, heading_mode: str) -> Dict[str, Any]:
        roots = list(self.roots)
        # 过滤掉空的初始前言节点（如果它既无正文也无表格且有其他有效大纲）
        if len(roots) > 1 and not roots[0].paragraphs and not roots[0].tables:
            roots.pop(0)
        return {
            "doc_path": doc_path,
            "total_sections": self.counter,
            "heading_mode": heading_mode,
            "sections": [sec.to_dict() for sec in roots],
            "tables": self.tables,
            "full_text": "\n".join(self.text_lines),
        }


class _StyleResolver:
    """
    单次解析内的样式查找缓存。python-docx 的 paragraph.style 每次调用都线性扫描全部样式，
    真实招标文件（~600 个样式 × 数千段落）单份解析耗时十余秒；按 style_id 缓存后降到秒内。
    每次 parse_docx 新建实例（解析器单例被多个后台任务线程共享，不能挂在实例属性上）。
    """

    def __init__(self, doc):
        self._by_id = {}
        self._default_id = None
        for el in doc.styles.element.findall(qn("w:style")):
            sid = el.get(qn("w:styleId"))
            self._by_id[sid] = el
            if el.get(qn("w:type")) == "paragraph" and el.get(qn("w:default")) in ("1", "true", "on"):
                self._default_id = sid
        self._levels = {}

    def style_id_of(self, para: Paragraph) -> Optional[str]:
        pPr = para._p.pPr
        if pPr is not None and pPr.pStyle is not None:
            return pPr.pStyle.val
        return self._default_id

    def name(self, style_id: Optional[str]) -> str:
        el = self._by_id.get(style_id)
        if el is None:
            return ""
        name_el = el.find(qn("w:name"))
        return (name_el.get(qn("w:val")) if name_el is not None else "").strip().lower()

    def heading_level(self, style_id: Optional[str], parser: "WordDocumentParser") -> Optional[int]:
        """样式名（heading/标题 N、title）或样式继承链上的大纲级别，按 style_id 缓存"""
        if style_id in self._levels:
            return self._levels[style_id]
        level, sid, depth = None, style_id, 0
        while sid is not None and depth < 8:
            name = self.name(sid)
            if name == "title":
                level = 1
                break
            m = parser.STYLE_HEADING.match(name)
            if m:
                level = max(1, int(m.group(1)))
                break
            el = self._by_id.get(sid)
            if el is None:
                break
            level = parser._outline_level_of(el.find(qn("w:pPr")))
            if level:
                break
            based = el.find(qn("w:basedOn"))
            sid = based.get(qn("w:val")) if based is not None else None
            depth += 1
        self._levels[style_id] = level
        return level


class WordDocumentParser:
    """
    专门针对中文招投标文件 / 技术标书的层级化 Word (.docx) 解析器：
    1. 标题识别"样式优先"：文档使用了标题样式（Heading/标题 N 或大纲级别）时只信任样式，
       正文里的"1. xxx"列表项不会被误判为标题（真实招标文件实测：正则模式把 67 个标题切成 690 个节点）；
       文档未使用标题样式时，才退回序号正则（第一章 / 1.1 / 一、）启发式识别
    2. 跳过目录（toc 样式 / 制表符+页码结尾的目录行），递归进入内容控件（w:sdt）
    3. 维护完整大纲层级树与面包屑路径（Breadcrumb）
    4. 表格去除横向合并单元格造成的重复，并记录所在章节与表前文字（供评分表等结构化抽取）
    5. 输出按文档顺序的全文 full_text（标题 + 正文 + 表格行），供拆标/偏离表/合规使用
    6. 层级归一：真实文件常有同一编号体例的标题样式不一致（如「第二篇」用正文样式 + 大纲 2 级、
       「十、」被直接设成 1 级），按编号体例（第X篇 / 一、 / （一） / 1. / 1.1 …）取文档内多数层级
    """

    # 正则识别标题层级（仅用于未使用标题样式的文档）
    REGEX_H1 = [
        re.compile(r"^第[一二三四五六七八九十百0-9]+[章节篇部]\s*.*"),
        re.compile(r"^[一二三四五六七八九十]+[、.．]\s*.*"),
        re.compile(r"^[0-9]{1,2}\s+[^\d\s\W].*"),
        re.compile(r"^[0-9]{1,2}[、.．]\s*.*"),
    ]
    REGEX_H2 = [
        re.compile(r"^[0-9]{1,2}\.[0-9]{1,2}\s*.*"),
        re.compile(r"^[（(][一二三四五六七八九十]+[）)]\s*.*"),
    ]
    REGEX_H3 = [
        re.compile(r"^[0-9]{1,2}\.[0-9]{1,2}\.[0-9]{1,2}\s*.*"),
        re.compile(r"^[0-9]{1,2}[）)]\s*.*"),
        re.compile(r"^[（(][0-9]{1,2}[）)]\s*.*"),
    ]
    # 目录行：标题 + 制表符/引导符 + 页码（如 "第一篇 投标邀请书\t- 3 -"、"1.1 概述……12"）
    TOC_LINE = re.compile(r"(?:\t|[.…·]{3,})\s*[-—]?\s*\d{1,4}\s*[-—]?\s*$")
    STYLE_HEADING = re.compile(r"^(?:heading|标题)\s*(\d)$")
    # 启用样式模式所需的最少标题段落数
    MIN_STYLED_HEADINGS = 3
    MAX_STYLED_HEADING_CHARS = 50
    SENTENCE_END = "。；;！!？?：:，,"
    # 「联系人：周老师」「（一）采购人：XX学院」这类"短标签：值"行即便带大纲级别也是正文
    LABEL_VALUE = re.compile(r"^[^：:]{1,8}[：:]\s*\S")
    NUMBER_CLASSES = [
        ("part", re.compile(r"^第[一二三四五六七八九十百\d]+[篇部章]")),
        ("cn", re.compile(r"^[一二三四五六七八九十]+、")),
        ("cn_paren", re.compile(r"^[（(][一二三四五六七八九十]+[）)]")),
        ("num3", re.compile(r"^\d{1,2}\.\d{1,2}\.\d{1,2}")),
        ("num2", re.compile(r"^\d{1,2}\.\d{1,2}(?![.\d])")),
        ("num1", re.compile(r"^\d{1,2}[.、．](?!\d)")),
        ("num_paren", re.compile(r"^[（(]\d{1,2}[）)]")),
    ]

    def __init__(self):
        pass

    # ---------------- 标题识别 ----------------

    @staticmethod
    def _outline_level_of(pPr) -> Optional[int]:
        if pPr is None:
            return None
        node = pPr.find(qn("w:outlineLvl"))
        if node is None:
            return None
        try:
            val = int(node.get(qn("w:val")))
        except (TypeError, ValueError):
            return None
        return val + 1 if 0 <= val <= 8 else None  # 9 = 正文级别

    def _style_heading_level(self, para: Paragraph, styles: "_StyleResolver") -> Optional[int]:
        """通过段落大纲级别 / 样式名 / 样式继承链上的大纲级别判断标题层级"""
        level = self._outline_level_of(para._p.pPr)
        if level:
            return level
        return styles.heading_level(styles.style_id_of(para), self)

    def _regex_heading_level(self, text: str) -> Optional[int]:
        # 标题一般不以逗号、分号或顿号结尾，且长度一般在 2~60 字符
        if len(text) > 60 or len(text) < 2 or text[-1] in ["，", ",", "；", ";", "、"]:
            return None
        # 过滤年份、日期或金额假标题（如 "2026 年是...", "100 万元"）
        if re.match(r"^(?:19|20)\d\d\s*年", text) or re.match(r"^\d+\s*(?:万|千|亿|元|个|条|家|台|套)", text):
            return None
        for pat in self.REGEX_H3:
            if pat.match(text):
                return 3
        for pat in self.REGEX_H2:
            if pat.match(text):
                return 2
        for pat in self.REGEX_H1:
            if pat.match(text):
                return 1
        return None

    def _detect_heading_level(self, para: Paragraph, styles: "_StyleResolver", use_styles: bool = True) -> Optional[int]:
        text = para.text.strip()
        if not text:
            return None
        if use_styles:
            level = self._style_heading_level(para, styles)
            # 招标文件常把整句条款也套上标题样式：成句的（过长或以句读结尾）按正文处理
            if (not level or len(text) > self.MAX_STYLED_HEADING_CHARS or text[-1] in self.SENTENCE_END
                    or self.LABEL_VALUE.match(text)):
                return None
            return min(level, 6)
        return self._regex_heading_level(text)

    def _number_class(self, text: str) -> Optional[str]:
        stripped = text.lstrip("※★▲◆ ")
        for name, pattern in self.NUMBER_CLASSES:
            if pattern.match(stripped):
                return name
        return None

    def _class_levels(self, headings: List[tuple]) -> Dict[str, int]:
        """每种编号体例在文档内的多数层级（≥2 个样本才生效，平票取较高层级）"""
        votes: Dict[str, Counter] = {}
        for text, level in headings:
            cls = self._number_class(text)
            if cls:
                votes.setdefault(cls, Counter())[level] += 1
        result = {}
        for cls, counter in votes.items():
            if sum(counter.values()) >= 2:
                top = max(counter.values())
                result[cls] = min(lvl for lvl, n in counter.items() if n == top)
        return result

    def _is_toc_paragraph(self, para: Paragraph, styles: "_StyleResolver") -> bool:
        name = styles.name(styles.style_id_of(para))
        if name.startswith("toc") or name.startswith("目录"):
            return True
        # 无目录样式时：页码前的文字须形似标题，避免误删"合计\t100"这类正文
        m = self.TOC_LINE.search(para.text)
        return bool(m) and self._regex_heading_level(para.text[:m.start()].strip()) is not None

    # ---------------- 文档遍历 ----------------

    @staticmethod
    def _iter_block_elements(container):
        """按文档顺序产出段落与表格元素；递归进入内容控件 w:sdt"""
        for el in container:
            if el.tag in (qn("w:p"), qn("w:tbl")):
                yield el
            elif el.tag == qn("w:sdt"):
                content = el.find(qn("w:sdtContent"))
                if content is not None:
                    yield from WordDocumentParser._iter_block_elements(content)

    def _extract_table(self, table: Table, cell_keys: Optional[List[List[int]]] = None) -> List[List[str]]:
        """
        提取表格文本。cell_keys 若传入，则同步填充每个单元格的物理单元格编号：
        纵向合并的单元格在各行编号相同，供评分表区分"一个评分项跨多行"与"相邻多项同分值"。
        """
        table_data = []
        tc_ids: Dict[int, int] = {}
        alive = []  # 持有 lxml 元素代理的强引用，保证 id() 在遍历期间稳定且不被复用
        for row in table.rows:
            cleaned_row, row_keys = [], []
            prev_tc = None
            for cell in row.cells:
                # 横向合并（gridSpan）的单元格会被重复返回，只保留一次；纵向合并保留（分类列逐行可读）
                if cell._tc is prev_tc:
                    continue
                prev_tc = cell._tc
                alive.append(prev_tc)
                cleaned_row.append(cell.text.strip().replace("\n", " "))
                row_keys.append(tc_ids.setdefault(id(prev_tc), len(tc_ids)))
            if any(cleaned_row):
                table_data.append(cleaned_row)
                if cell_keys is not None:
                    cell_keys.append(row_keys)
        return table_data

    def parse_docx(self, file_path: str or Path) -> Dict[str, Any]:
        doc = docx.Document(str(file_path))
        styles = _StyleResolver(doc)
        elements = list(self._iter_block_elements(doc.element.body))

        # 第一遍：统计标题样式段落数决定识别模式，并收集各编号体例的层级用于归一
        styled = 0
        candidates: List[tuple] = []
        for el in elements:
            if el.tag == qn("w:p"):
                para = Paragraph(el, doc)
                if para.text.strip() and not self._is_toc_paragraph(para, styles):
                    if self._style_heading_level(para, styles):
                        styled += 1
                    level = self._detect_heading_level(para, styles, use_styles=True)
                    if level:
                        candidates.append((para.text.strip(), level))
        use_styles = styled >= self.MIN_STYLED_HEADINGS
        class_levels = self._class_levels(candidates) if use_styles else {}

        builder = DocumentTreeBuilder()
        for el in elements:
            if el.tag == qn("w:p"):
                para = Paragraph(el, doc)
                text = para.text.strip()
                if not text or self._is_toc_paragraph(para, styles):
                    continue
                level = self._detect_heading_level(para, styles, use_styles)
                if level is not None:
                    builder.add_heading(text, class_levels.get(self._number_class(text), level))
                else:
                    builder.add_paragraph(text)
            else:
                keys: List[List[int]] = []
                tbl_data = self._extract_table(Table(el, doc), cell_keys=keys)
                builder.add_table(tbl_data, keys)

        return builder.result(str(file_path), "styles" if use_styles else "regex")

    def flatten_sections_to_chunks(self, parsed_result: Dict[str, Any], doc_name: str) -> List[Dict[str, Any]]:
        """
        将大纲树扁平化为可直接入库的知识切片（Knowledge Chunks），每个切片携带完整面包屑大纲路径
        """
        chunks = []

        def traverse(node: Dict[str, Any]):
            content = node.get("content", "")
            tables = node.get("tables", [])

            # 如果包含表格，将表格转换成 Markdown 格式拼入正文
            table_md_list = []
            for tbl in tables:
                if not tbl:
                    continue
                header = tbl[0]
                sep = ["---"] * len(header)
                md_rows = ["| " + " | ".join(header) + " |", "| " + " | ".join(sep) + " |"]
                for row in tbl[1:]:
                    # 补齐长度
                    padded = row + [""] * (len(header) - len(row))
                    md_rows.append("| " + " | ".join(padded[:len(header)]) + " |")
                table_md_list.append("\n".join(md_rows))

            full_text = content
            if table_md_list:
                full_text += "\n\n" + "\n\n".join(table_md_list)

            if full_text.strip():
                # 自动提炼标签（根据标题和关键词规则）
                tags = []
                keywords_map = {
                    "架构": ["系统架构", "技术路线"],
                    "安全": ["安全方案", "等保", "数据安全"],
                    "运维": ["售后服务", "运维保障", "SLA"],
                    "实施": ["项目实施", "进度管理", "部署方案"],
                    "测试": ["软件测试", "压力测试", "验收标准"],
                    "培训": ["人员培训", "交付培训"],
                    "信创": ["信创适配", "国产化"],
                    "数据库": ["数据库设计", "容灾备份"],
                    "微服务": ["微服务", "容器化", "Kubernetes"]
                }
                for kw, mapped_tags in keywords_map.items():
                    if kw in node["breadcrumb"] or kw in full_text[:100]:
                        tags.extend(mapped_tags)

                chunks.append({
                    "id": f"{doc_name}_{node['section_id']}",
                    "doc_name": doc_name,
                    "section_title": node["title"],
                    "breadcrumb": node["breadcrumb"],
                    "content": full_text.strip(),
                    "level": node["level"],
                    "tags": list(set(tags))
                })

            for sub in node.get("subsections", []):
                traverse(sub)

        for sec in parsed_result.get("sections", []):
            traverse(sec)

        return chunks
