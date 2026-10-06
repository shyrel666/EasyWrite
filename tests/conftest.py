"""
测试公共设施：临时运行时数据、路径注入、样例文档构造与项目创建辅助。
"""
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest

BASE_DIR = Path(__file__).resolve().parent.parent / "backend"
sys.path.insert(0, str(BASE_DIR))

# 必须在导入应用之前隔离：数据库和配置管理器会在模块导入时初始化。
_runtime_dir = TemporaryDirectory(prefix="easywrite-tests-")
TEST_DATA_DIR = Path(_runtime_dir.name).resolve()
_test_env = pytest.MonkeyPatch()
for name, relative_path in {
    "DATA_DIR": "",
    "DB_PATH": "easywrite.db",
    "UPLOAD_DIR": "uploads",
    "KNOWLEDGE_DIR": "knowledge_store",
    "OUTPUT_DIR": "exports",
    "RENDERED_DIR": "rendered_diagrams",
    "TEMPLATE_DIR": "templates",
}.items():
    _test_env.setenv(name, str(TEST_DATA_DIR / relative_path))

# 不读取用户密钥；环境变量的空值也会覆盖本地 .env 中的配置。
_test_env.setenv("LLM_API_KEY", "")
_test_env.setenv("EMBEDDING_API_KEY", "")

from fastapi.testclient import TestClient  # noqa: E402
from app.main import app  # noqa: E402

TENDER_SAMPLE = """项目名称：智慧水务一体化综合调度系统建设项目
项目编号：SW-2026-ZB-088
采购人：某市水务环境集团有限公司
最高限价：850.00万元人民币
工期要求：合同签订后120个日历日内完成系统终验上线
质保期：免费提供5年驻场质保及7×24小时运维
投标保证金：17万元
投标有效期：自投标截止之日起90个日历日
评标办法：综合评分法。商务分20分，技术分50分，价格分30分。
★ 投标人所投核心调度引擎必须具备国家版权局颁发的独立软件著作权。
★ 核心数据层必须全面支持国产信创数据库（达梦或人大金仓），并支持国密SM4加密。
★ 承诺系统发生一级重大调度故障时，10分钟内响应，30分钟内工程师到达现场。
资质要求：具备CMMI5级认证或国家信息系统建设和服务能力CS3级。
实施地点：采购人指定项目现场
付款方式：合同签订后支付30%，初验支付30%，终验支付30%，质保期满支付10%。
投标截止时间：2026年9月30日09:30
"""


def pytest_unconfigure():
    """收集测试或执行测试后，先释放 SQLite 连接，再清理临时目录。"""
    from app.db.database import engine

    engine.dispose()
    _test_env.undo()
    _runtime_dir.cleanup()


@pytest.fixture(scope="session")
def runtime_data_dir():
    return TEST_DATA_DIR


@pytest.fixture(autouse=True)
def isolate_runtime_config():
    """
    在临时目录中逐用例快照并恢复配置，避免假 Key 或资产修改污染后续用例。
    """
    data_dir = TEST_DATA_DIR
    files = ["ai_settings.json", "templates.json", "enterprise_assets.json"]
    snapshots = {}
    for name in files:
        path = data_dir / name
        snapshots[name] = path.read_bytes() if path.exists() else None
    yield
    for name, content in snapshots.items():
        path = data_dir / name
        if content is None:
            path.unlink(missing_ok=True)
        else:
            path.write_bytes(content)
    # 重载内存态配置与客户端
    from app.core.ai_settings_manager import ai_settings_manager
    from app.core.llm_client import llm_client
    from app.core.embedding_client import embedding_client
    from app.services.assets.asset_manager import asset_manager
    from app.services.exporter.template_manager import template_manager
    ai_settings_manager._load()
    llm_client.reload_config()
    embedding_client.reload_config()
    asset_manager._load_or_init()
    template_manager.templates.clear()
    template_manager._init_templates()


@pytest.fixture(scope="session")
def client():
    return TestClient(app)


@pytest.fixture()
def project_id(client):
    """创建并返回一个测试项目 id（测试结束后清理）"""
    res = client.post("/api/v1/project/create", json={
        "name": "测试项目_智慧水务调度系统",
        "client_name": "某市水务环境集团",
        "description": "物联感知、实时调度、能耗管理与安防联动一体化平台",
    })
    assert res.status_code == 200
    pid = res.json()["id"]
    yield pid
    client.delete(f"/api/v1/project/{pid}")


def build_sample_bid_docx(path: Path) -> Path:
    """构造一份带层级大纲与表格的样例历史标书"""
    import docx as docxlib
    doc = docxlib.Document()
    doc.add_heading("第一章 总体技术方案", level=1)
    doc.add_paragraph("本系统采用云原生微服务架构，基于 Kubernetes 容器编排实现服务注册发现、弹性伸缩与灰度发布。")
    doc.add_paragraph("系统总体分为接入层、网关层、微服务业务层与数据持久化层，各层通过 RESTful API 与消息队列解耦。")
    doc.add_heading("1.1 微服务治理设计", level=2)
    doc.add_paragraph("服务治理采用 Spring Cloud Alibaba 体系，包括 Nacos 注册中心、Sentinel 限流熔断与 Seata 分布式事务。")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "组件"
    table.cell(0, 1).text = "用途"
    table.cell(0, 2).text = "选型"
    table.cell(1, 0).text = "注册中心"
    table.cell(1, 1).text = "服务发现"
    table.cell(1, 2).text = "Nacos"
    doc.add_heading("第二章 数据安全与等保设计", level=1)
    doc.add_paragraph("系统按照国家网络安全等级保护三级标准设计，采用国密 SM2/SM3/SM4 算法对核心数据加密存储。")
    doc.add_paragraph("建立完整的数据库审计与入侵检测机制，核心业务数据每日全量备份、每小时增量备份。")
    doc.add_heading("第三章 售后运维保障方案", level=1)
    doc.add_paragraph("提供7×24小时运维支持，5分钟内响应，30分钟内远程介入，重大故障2小时内工程师到达现场。")
    doc.save(str(path))
    return path
