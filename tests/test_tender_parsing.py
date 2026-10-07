"""
招标文件解析与评分细则抽取：
- 解析器：目录跳过、样式优先标题识别、成句标题按正文、全文保真
- 评分表：纵向合并续行 vs 相邻同分值、分值写法变体、分包识别、非评分表排除
- 大模型 18 项抽取路径回归（model_fields 误用曾导致 LLM 抽取永远失败）
- 上传接口：评分细则随任务返回、招标原文存档到项目
- 真实招标文件（downloads/ccgp_* 存在时）：每包分值合计 100
"""
import time
from functools import lru_cache
from pathlib import Path

import docx as docxlib
import pytest
from docx.enum.style import WD_STYLE_TYPE

from app.core.llm_client import llm_client
from app.services.parser.scoring_extractor import extract_scoring_items, parse_points_cell, summarize
from app.services.parser.tender_analyzer import tender_analyzer
from app.services.parser.word_parser import WordDocumentParser

REAL_DIR = Path(__file__).resolve().parent.parent / "downloads" / "ccgp_2026-10-06"
HEADER = ["序号", "评分因素及权重", "分值", "评分标准", "说明"]


def _fill_rubric(table, rows, merges):
    for r, values in enumerate(rows):
        for c, text in enumerate(values):
            table.cell(r, c).text = text
    for (r1, c1), (r2, c2), text in merges:
        table.cell(r1, c1).merge(table.cell(r2, c2)).text = text


def build_tender_docx(path: Path) -> Path:
    doc = docxlib.Document()
    toc = doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("第一篇 投标邀请书\t- 3 -", style=toc)
    doc.add_paragraph("第四篇 评标办法\t- 20 -", style=toc)
    doc.add_heading("第一篇 投标邀请书", level=1)
    doc.add_paragraph("项目名称：智慧园区运维服务项目")
    doc.add_heading("一、服务需求", level=2)
    doc.add_heading("1.服务范围：为采购人提供网络安全服务。", level=3)  # 成句的标题样式 → 正文
    doc.add_paragraph("1.投标人具有ISO20000证书得1分")  # 正文编号列表，不是标题
    doc.add_paragraph("合计\t100")
    doc.add_heading("第四篇 评标办法", level=1)

    exam = doc.add_table(rows=2, cols=4)  # 履约考核表：不是评分表
    _fill_rubric(exam, [["序号", "考核项目", "考核内容", "扣分标准"], ["1", "系统崩溃", "核心功能不可用", "每次扣2分"]], [])
    qual = doc.add_table(rows=2, cols=3)  # 符合性审查表：无分值列
    _fill_rubric(qual, [["序号", "评审因素", "评审标准"], ["1", "投标文件签署", "签署齐全"]], [])

    doc.add_paragraph("（一）包1评审因素")
    t1 = doc.add_table(rows=8, cols=5)
    _fill_rubric(t1, [
        HEADER,
        ["1", "投标报价（30%）", "30", "有效的投标报价中的最低价为评标基准价，其价格分为满分。", ""],
        ["2", "", "技术方案（30分）", "1.总体设计（10分）方案完整得10分。2.实施方案（20分）方案可行得20分。", "格式自拟"],
        ["2", "", "10分", "1. 项目理解 投标人对本项目的理解阐述。", ""],
        ["2", "", "10分", "2. 运维工作方案 投标人提供运维工作方案。", ""],
        ["3", "", "人员要求（12分）", "1.项目经理具备高级职称得6分。", "提供证书"],
        ["3", "", "", "2.技术团队每具备一个证书得2分，最高6分。", "提供证书"],
        ["3", "", "业绩8", "每提供一个类似业绩得2分，最高8分。", "提供合同"],
    ], [
        ((2, 1), (4, 1), "技术部分（50%）"),
        ((5, 1), (7, 1), "商务部分（20%）"),
        ((5, 2), (6, 2), "人员要求（12分）"),  # 一个评分项跨两行
    ])

    doc.add_paragraph("（二）包2评审因素")
    t2 = doc.add_table(rows=3, cols=5)
    _fill_rubric(t2, [
        HEADER,
        ["1", "投标报价（40%）", "40分", "有效的投标报价中的最低价为评标基准价。", ""],
        ["2", "技术部分（60%）", "现场 演示60分", "投标人现场演示系统功能。", ""],
    ], [])
    doc.save(str(path))
    return path


@pytest.fixture(scope="module")
def tender_docx(tmp_path_factory):
    return build_tender_docx(tmp_path_factory.mktemp("tender") / "tender.docx")


@pytest.fixture(scope="module")
def parsed(tender_docx):
    return WordDocumentParser().parse_docx(tender_docx)


def _titles(sections):
    out = []
    for s in sections:
        out.append(s["title"])
        out.extend(_titles(s["subsections"]))
    return out


def test_parser_skips_toc_and_trusts_heading_styles(parsed):
    titles = _titles(parsed["sections"])
    assert parsed["heading_mode"] == "styles"
    assert titles.count("第一篇 投标邀请书") == 1, "目录行不应成为标题"
    assert not any("\t" in t for t in titles)
    assert "1.投标人具有ISO20000证书得1分" not in titles, "正文编号列表被误判为标题"
    assert "1.服务范围：为采购人提供网络安全服务。" not in titles, "成句的标题样式应按正文处理"
    full = parsed["full_text"]
    assert "1.服务范围：为采购人提供网络安全服务。" in full
    assert "合计\t100" in full, "非标题形态的制表符行不能当目录删除"
    assert "技术方案（30分）" in full, "表格内容应进入全文"


def test_scoring_items_from_merged_tables(parsed):
    items = extract_scoring_items(parsed["tables"])
    pkg1 = [it for it in items if it.package == "包1"]
    pkg2 = [it for it in items if it.package == "包2"]
    assert [it.name for it in pkg1] == ["投标报价", "技术方案", "项目理解", "运维工作方案", "人员要求", "业绩"]
    assert [it.points for it in pkg1] == [30, 30, 10, 10, 12, 8]
    assert sum(it.points for it in pkg1) == 100

    staff = pkg1[4]
    assert "项目经理" in staff.criteria and "技术团队" in staff.criteria, "跨行评分项应合并评分标准"
    assert staff.category == "商务部分" and staff.category_weight == 20
    assert staff.response_type == "evidence"

    plan = pkg1[1]
    assert [(s.name, s.points) for s in plan.sub_items] == [("总体设计", 10), ("实施方案", 20)]
    assert plan.kind == "technical" and plan.response_type == "proposal" and plan.note == "格式自拟"
    assert pkg1[0].kind == "price"

    assert [(it.name, it.points, it.response_type) for it in pkg2] == [
        ("投标报价", 40, "price"), ("现场演示", 60, "demo"),
    ]
    assert summarize(items) == "包1：6 项，合计 100 分；包2：2 项，合计 100 分"


@pytest.mark.parametrize("cell,expected", [
    ("技术服务方案 （50分）", ("技术服务方案", 50)),
    ("20分", ("", 20)),
    ("30", ("", 30)),
    ("基础服务20", ("基础服务", 20)),
    ("服务 方案 12", ("服务方案", 12)),
    ("现场 演示24分", ("现场演示", 24)),
    ("4.5分", ("", 4.5)),
    ("报价 （20分）", ("报价", 20)),
    ("见评分标准", ("见评分标准", None)),
])
def test_parse_points_cell_variants(cell, expected):
    assert parse_points_cell(cell) == expected


def test_summarize_flags_inconsistent_total(parsed):
    items = extract_scoring_items(parsed["tables"])
    items[0].points = 25
    assert "与满分 100 不一致" in summarize(items)
    assert "未在招标文件中识别到评分表" in summarize([])


def test_llm_tender_extraction_path_is_used(monkeypatch):
    """回归：TenderAnalysis18.model_fields() 误当方法调用，导致配置模型后拆标仍永远走规则兜底"""
    monkeypatch.setattr(llm_client, "is_configured", True)
    monkeypatch.setattr(llm_client, "get_mode", lambda: "llm")
    monkeypatch.setattr(llm_client, "chat_completion_structured", lambda **kw: {
        "project_name": "大模型抽取的项目名", "star_disqualification_items": ["★ 必须提供原厂授权"],
    })
    result = tender_analyzer.analyze_text("项目名称：规则抽取的项目名\n★ 必须提供原厂授权")
    assert result.extraction_mode == "llm"
    assert result.project_name == "大模型抽取的项目名"


def test_upload_analysis_returns_scoring_and_archives_text(client, project_id, tender_docx):
    with open(tender_docx, "rb") as f:
        res = client.post("/api/v1/tender/analyze", data={"project_id": project_id},
                          files={"file": ("tender.docx", f, "application/octet-stream")})
    assert res.status_code == 200
    task_id = res.json()["task_id"]
    for _ in range(100):
        task = client.get(f"/api/v1/tasks/{task_id}").json()
        if task["status"] in ("completed", "failed"):
            break
        time.sleep(0.05)
    assert task["status"] == "completed", task.get("error")
    result = task["result"]
    assert len(result["scoring_items"]) == 8
    assert result["target_package"] == "包1"
    assert result["tech_score_weight"].startswith("50%")
    text = client.get(f"/api/v1/project/{project_id}/tender/text").json()
    assert text["has_text"] and "评分标准" in text["text"]


def test_upload_with_unknown_project_rejected(client, tender_docx):
    with open(tender_docx, "rb") as f:
        res = client.post("/api/v1/tender/analyze", data={"project_id": "proj_missing"},
                          files={"file": ("tender.docx", f, "application/octet-stream")})
    assert res.status_code == 404


# ---------------- 真实招标文件（仓库不附带，存在时验证） ----------------

REAL_EXPECTED = {
    "01_": {"包1": 6, "包2": 7},
    "02_": {"": 7},
    "03_": {"": 11},
    "04_": {"包1": 6, "包2": 7},
    "05_": {"": 7},
}


def _real_file(prefix: str) -> Path:
    matches = sorted(REAL_DIR.glob(f"{prefix}*.docx")) if REAL_DIR.exists() else []
    if not matches:
        pytest.skip("未找到真实招标文件 downloads/ccgp_2026-10-06")
    return matches[0]


@lru_cache(maxsize=None)
def _parse_real(path: Path) -> dict:
    return WordDocumentParser().parse_docx(path)


@pytest.mark.parametrize("prefix", sorted(REAL_EXPECTED))
def test_real_tender_rubrics_total_100(prefix):
    parsed = _parse_real(_real_file(prefix))
    items = extract_scoring_items(parsed["tables"])
    counts = {}
    for it in items:
        counts[it.package] = counts.get(it.package, 0) + 1
    assert counts == REAL_EXPECTED[prefix]
    for pkg in counts:
        assert sum(it.points for it in items if it.package == pkg) == pytest.approx(100)


@pytest.mark.parametrize("prefix", sorted(REAL_EXPECTED))
def test_real_tender_full_text_retained(prefix):
    path = _real_file(prefix)
    raw = sum(len(t.text or "") for t in docxlib.Document(str(path)).element.body.iter(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
    parsed = _parse_real(path)
    kept = len(parsed["full_text"].replace(" | ", "").replace("\n", ""))
    assert kept / raw > 0.95
    assert not any(WordDocumentParser.TOC_LINE.search(t) for t in _titles(parsed["sections"])), "目录行不应成为标题"


# ---------------- 条款标记含义（以文件自身定义为准） ----------------

LEGEND_TEXT = """注：本节“※”标注的技术需求为符合性审查中的实质性要求，若不满足按无效投标处理。
本节“★”标注的技术需求为重要技术需求，若不满足将按照评标因素中相关规定处理。
※1.服务期内须提供7×24小时驻场运维。
★2.支持国产化数据库达梦。
3 | 技术部分 | 投标文件内容 | 本招标文件第二篇中（※）号标注的部分。
4.一般性要求：按时提交月报。"""


def test_marker_legend_decides_red_lines():
    result = tender_analyzer.analyze_text(LEGEND_TEXT)
    assert result.star_disqualification_items == ["※1.服务期内须提供7×24小时驻场运维。"]
    assert result.important_items == ["★2.支持国产化数据库达梦。"], "按评分扣分的★条款不是废标红线"
    assert set(result.marker_legend) == {"※", "★"}
    assert "无效投标" in result.marker_legend["※"]


def test_marker_defaults_without_legend():
    result = tender_analyzer.analyze_text("★ 必须提供原厂授权书。\n▲ 支持双活部署。\n普通要求。")
    assert result.star_disqualification_items == ["★ 必须提供原厂授权书。"]
    assert result.important_items == ["▲ 支持双活部署。"]
    assert result.marker_legend == {}


def test_llm_cannot_override_rule_only_fields(monkeypatch):
    monkeypatch.setattr(llm_client, "is_configured", True)
    monkeypatch.setattr(llm_client, "get_mode", lambda: "llm")
    monkeypatch.setattr(llm_client, "chat_completion_structured", lambda **kw: {
        "project_name": "X", "marker_legend": "非法类型", "scoring_items": "非法",
        "important_items": ["★2.支持国产化数据库达梦。"],
    })
    result = tender_analyzer.analyze_text(LEGEND_TEXT)
    assert result.extraction_mode == "llm" and result.project_name == "X"
    assert "※" in result.marker_legend and result.scoring_items == []
    assert result.important_items == ["★2.支持国产化数据库达梦。"], "规则补充不应与模型结果重复"
