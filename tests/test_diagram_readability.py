"""逻辑拆图、跨图边保留、版面恢复和失败插槽。"""
import io
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from PIL import Image

from app.models.schemas import OutlineNode
from app.services.exporter.diagram_spec import diagram_parts, parse_spec
from app.services.exporter.docx_generator import docx_exporter


def graph(size):
    lines = ['graph TD']
    for start in range(0, size, 5):
        lines.append(f'subgraph 逻辑层{start // 5 + 1}')
        lines += [f'N{i}[业务节点{i + 1}]' for i in range(start, min(size, start + 5))]
        lines.append('end')
    lines += [f'N{i} --> N{i+1}' for i in range(size - 1)]
    return '\n'.join(lines)


def test_dense_parts_keep_all_nodes_and_cross_edges():
    for size in (30, 60):
        spec = parse_spec(graph(size))
        parts = diagram_parts(graph(size))
        assert len(parts) == size // 5
        assert set().union(*(parse_spec(p['code']).nodes.keys() for p in parts)) == set(spec.nodes)
        for a, b, _ in spec.edges:
            assert any([a, b, ''] in parse_spec(p['code']).edges or f'{spec.nodes[a]} → {spec.nodes[b]}' in p['connections'] for p in parts)


def test_wide_page_restores_portrait_and_failure_slot():
    png = io.BytesIO()
    Image.new('RGB', (1800, 500), 'white').save(png, format='PNG')
    code = 'graph LR\n A[入口] --> B[服务]'
    n = OutlineNode(id='s', title='网络', content=f'```mermaid\n{code}\n```\n后续纵向正文\n```mermaid\ninvalid diagram\n```')
    path = docx_exporter.export_project_to_docx('版面', '采购人', [n], diagrams={code: png.getvalue()})
    try:
        doc = Document(path)
        assert [s.orientation for s in doc.sections] == [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT]
        assert len(doc.inline_shapes) == 1
        assert '图 1' in '\n'.join(p.text for p in doc.paragraphs)
        assert '图 2' in '\n'.join(p.text for p in doc.paragraphs)
        assert any('未能自动渲染' in c.text for t in doc.tables for r in t.rows for c in r.cells)
        shape = doc.inline_shapes[0]
        assert abs(shape.width / shape.height - 3.6) < .01
    finally:
        path.unlink(missing_ok=True)


def test_typical_density_documents_are_renderable():
    for size in (10, 30, 60):
        n = OutlineNode(id='s', title=f'{size} 节点架构图', content=f'```mermaid\n{graph(size)}\n```')
        path = docx_exporter.export_project_to_docx('图表可读性回归', '采购人', [n])
        try:
            doc = Document(path)
            assert len(doc.inline_shapes) == (1 if size == 10 else size // 5)
            if os.getenv('EASYWRITE_QA_DIR'):
                (Path(os.environ['EASYWRITE_QA_DIR']) / f'p1-6-{size}.docx').write_bytes(path.read_bytes())
        finally:
            path.unlink(missing_ok=True)
