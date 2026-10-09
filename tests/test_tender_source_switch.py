"""R05：文件与粘贴文本切换时，正文、章节树、来源及应用的拆标结果保持一致。"""
from io import BytesIO
import time

from docx import Document
import pytest

from app.services.project_store import project_store


def _upload(client, project_id, marker):
    doc = Document()
    doc.add_heading("第一章 项目概况", level=1)
    doc.add_paragraph(f"项目名称：{marker}建设项目")
    doc.add_heading("第二章 项目技术需求", level=1)
    doc.add_paragraph(f"系统必须支持{marker}数据库备份，并提供完整的数据恢复功能。")
    data = BytesIO()
    doc.save(data)
    res = client.post("/api/v1/tender/analyze", data={"project_id": project_id},
                      files={"file": ("招标文件.docx", data.getvalue(),
                                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res.status_code == 200
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        task = client.get(f"/api/v1/tasks/{res.json()['task_id']}").json()
        if task["status"] in ("completed", "failed"):
            break
        time.sleep(0.02)
    assert task["status"] == "completed", task.get("error")
    assert client.post(f"/api/v1/project/{project_id}/tender/apply", json={"analysis": task["result"]}).status_code == 200
    return task["result"]


def _clauses(client, project_id):
    res = client.post(f"/api/v1/project/{project_id}/deviation/extract", json={})
    assert res.status_code == 200
    return "\n".join(item["clause_title"] for item in res.json()["items"])


def test_upload_then_paste_replaces_original_tree_and_can_switch_back(client, project_id):
    _upload(client, project_id, "旧系统甲")
    tree = client.get(f"/api/v1/project/{project_id}/tender/outline").json()
    assert tree["has_structure"]
    old_section = tree["sections"][1]["id"]
    assert "旧系统甲" in _clauses(client, project_id)

    text = "项目名称：新系统乙建设项目\n★ 系统必须支持新系统乙实时数据采集，并提供数据查询接口。"
    analysis = client.post("/api/v1/tender/analyze/text", json={"text": text}).json()
    assert client.post(f"/api/v1/project/{project_id}/tender/apply", json={
        "analysis": analysis, "tender_text": text}).status_code == 200
    stored = project_store.get_tender_document(project_id)
    assert stored == {"text": text, "structure": None, "source": "text"}
    assert client.get(f"/api/v1/project/{project_id}/tender/text").json()["text"] == text
    outline = client.get(f"/api/v1/project/{project_id}/tender/outline").json()
    assert outline["has_structure"] is False and outline["source"] == "text"
    assert client.get(f"/api/v1/project/{project_id}/tender/section", params={"path": old_section}).status_code == 404
    clauses = _clauses(client, project_id)
    assert "新系统乙" in clauses and "旧系统甲" not in clauses

    _upload(client, project_id, "新文件丙")
    stored = project_store.get_tender_document(project_id)
    assert stored["source"] == "docx" and stored["structure"]["sections"]
    clauses = _clauses(client, project_id)
    assert "新文件丙" in clauses and "新系统乙" not in clauses and "旧系统甲" not in clauses


def test_apply_rolls_back_analysis_and_original_document_together(client, project_id, monkeypatch):
    old_analysis = _upload(client, project_id, "原始项目")
    before = project_store.get_tender_document(project_id)
    write = project_store.set_tender_document_in

    def broken(session, *args, **kwargs):
        write(session, *args, **kwargs)
        raise RuntimeError("模拟原文写入后、事务提交前失败")

    monkeypatch.setattr(project_store, "set_tender_document_in", broken)
    with pytest.raises(RuntimeError, match="事务提交前失败"):
        client.post(f"/api/v1/project/{project_id}/tender/apply", json={
            "analysis": {"project_name": "新项目"}, "tender_text": "系统必须支持新功能。"})
    assert project_store.get_tender_document(project_id) == before
    assert project_store.get(project_id).tender_analysis.model_dump() == old_analysis


def test_plain_text_store_write_also_clears_file_structure(client, project_id):
    project_store.set_tender_document(project_id, "文件正文", {"sections": [{"title": "旧章节"}], "source_format": "pdf"})
    project_store.set_tender_document(project_id, "粘贴正文")
    assert project_store.get_tender_document(project_id) == {"text": "粘贴正文", "structure": None, "source": "text"}
