"""原生评分索引：只导出目标包，按 ID 关联全部章节，不伪造页码。"""
import io
import os
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from app.models.schemas import OutlineNode, TenderAnalysis18
from app.services.exporter.docx_generator import docx_exporter
from app.services.exporter.scoring_index import scoring_index_rows
from app.services.project_store import project_store


def make_project(pid):
    p = project_store.get(pid)
    p.tender_analysis = TenderAnalysis18(target_package="包1", scoring_items=[
        {"id": "i1", "name": "服务方案", "points": 20, "package": "包1"},
        {"id": "i2", "name": "服务方案", "points": 10, "package": "包1"},
        {"id": "i3", "name": "其他分包秘密", "points": 90, "package": "包2"},
    ])
    p.outline = [OutlineNode(id="s1", title="第一章 服务方案", content="故障受理后按分级方案安排处置。", scoring_item_ids=["i1"]),
                 OutlineNode(id="s2", title="第二章 服务方案", content="巡检过程记录及闭环管理。", scoring_item_ids=["i1"])]
    return p


def test_native_index_bookmarks_fields_and_unmapped_items(project_id):
    p = make_project(project_id)
    rows = scoring_index_rows(p)
    assert [r["id"] for r in rows] == ["i1", "i2"]
    assert len(rows[0]["sections"]) == 2 and not rows[1]["sections"]
    path = docx_exporter.export_project_to_docx(p.name, p.client_name, p.outline, scoring_index=rows)
    try:
        doc = Document(path)
        table = doc.tables[0]
        assert [c.text for c in table.rows[0].cells] == ["评分项", "分值", "响应章节", "证明附件", "页码"]
        assert table.cell(2, 2).text == "待补"
        assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        with ZipFile(path) as z:
            xml = z.read("word/document.xml").decode()
        assert "其他分包秘密" not in xml and xml.count("PAGEREF") == 2
        marks = doc.element.xpath("//w:bookmarkStart")
        assert len(marks) == len({m.get(qn('w:name')) for m in marks}) == 4
        assert len(marks) == len({m.get(qn('w:id')) for m in marks})
        assert all(s["bookmark"] in xml for s in rows[0]["sections"])
        if os.getenv("EASYWRITE_QA_DIR"):
            target = Path(os.environ["EASYWRITE_QA_DIR"])
            target.mkdir(parents=True, exist_ok=True)
            (target / "p1-4.docx").write_bytes(path.read_bytes())
    finally:
        path.unlink(missing_ok=True)


def test_export_switch_and_preview(client, project_id):
    p = make_project(project_id)
    project_store.update(project_id, lambda current: (setattr(current, "outline", p.outline), setattr(current, "tender_analysis", p.tender_analysis)))
    url = f"/api/v1/project/{project_id}/export"
    assert len(client.get(url + "/preflight").json()["scoring_index"]) == 2
    for enabled in (False, True):
        response = client.post(url, json={"include_scoring_index": enabled})
        assert response.status_code == 200
        document = Document(io.BytesIO(response.content))
        assert bool(document.tables) == enabled
