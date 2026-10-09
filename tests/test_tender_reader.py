"""阶段 B4：按章节读取招标原文（章节树、分段读取、要点与相关章节）"""
import docx as docxlib

from app.services.parser.document_parser import parse_document
from app.services.project_store import project_store


def _tender(tmp_path):
    doc = docxlib.Document()
    doc.add_heading("第一章 项目概况", level=1)
    doc.add_paragraph("本项目为某市政务云平台年度运维服务。")
    doc.add_heading("第二章 技术需求", level=1)
    doc.add_heading("2.1 运维服务要求", level=2)
    doc.add_paragraph("投标人须提供7×24小时运维服务，并制定完善的应急预案，每季度开展一次应急演练。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "服务项", "要求"
    table.cell(1, 0).text, table.cell(1, 1).text = "巡检", "每月一次"
    doc.add_heading("2.2 人员要求", level=2)
    doc.add_paragraph("项目经理须具备信息系统项目管理师证书。" * 40)
    path = tmp_path / "tender.docx"
    doc.save(str(path))
    return parse_document(path)


def _setup(client, pid, tmp_path):
    parsed = _tender(tmp_path)
    project_store.set_tender_document(pid, parsed["full_text"], parsed)
    outline = [{"id": "sec_1", "title": "第一章 技术服务方案", "level": 1, "children": [
        {"id": "sec_1_1", "title": "1.1 应急预案", "level": 2,
         "requirements": ["评分要点：应急预案", "所属评分项：技术服务方案（30分）"]}]}]
    client.put(f"/api/v1/project/{pid}/outline", json={"outline": outline})
    return parsed


def _flat(nodes):
    for n in nodes:
        yield n
        yield from _flat(n["children"])


def test_outline_tree_and_related_sections(client, project_id, tmp_path):
    _setup(client, project_id, tmp_path)
    res = client.get(f"/api/v1/project/{project_id}/tender/outline").json()
    assert res["has_structure"]
    titles = [n["title"] for n in _flat(res["sections"])]
    assert titles == ["第一章 项目概况", "第二章 技术需求", "2.1 运维服务要求", "2.2 人员要求"]
    chapter2 = res["sections"][1]
    assert chapter2["level"] == 1 and chapter2["chars"] == 0
    assert chapter2["total_chars"] == sum(c["total_chars"] for c in chapter2["children"]) > 0
    assert chapter2["children"][0]["path"] == "第二章 技术需求 > 2.1 运维服务要求"

    with_node = client.get(f"/api/v1/project/{project_id}/tender/outline?section_id=sec_1_1").json()
    assert with_node["keywords"][:2] == ["应急预案", "技术服务方案"]
    assert [r["title"] for r in with_node["related"]] == ["2.1 运维服务要求"]
    assert with_node["related"][0]["hits"] == ["应急预案"]


def test_section_text_paging(client, project_id, tmp_path):
    _setup(client, project_id, tmp_path)
    tree = client.get(f"/api/v1/project/{project_id}/tender/outline").json()["sections"]
    chapter2 = tree[1]
    url = f"/api/v1/project/{project_id}/tender/section"

    full = client.get(url, params={"path": chapter2["id"]}).json()
    assert full["title"] == "第二章 技术需求" and not full["truncated"]
    assert "## 2.1 运维服务要求" in full["text"] and "巡检 | 每月一次" in full["text"] and "## 2.2 人员要求" in full["text"]
    own = client.get(url, params={"path": chapter2["id"], "deep": False}).json()
    assert own["text"] == "" and own["total"] == 0

    # 按完整路径读取 + 分段
    part = client.get(url, params={"path": "第二章 技术需求 > 2.2 人员要求", "limit": 100}).json()
    assert part["start"] == 0 and part["end"] == 100 and part["truncated"] and part["total"] > 100
    rest = client.get(url, params={"path": part["id"], "offset": part["end"], "limit": 100000}).json()
    assert rest["start"] == 100 and rest["end"] == rest["total"] and rest["truncated"]  # 从中间读起也标为非完整
    assert part["text"] + rest["text"] == client.get(url, params={"path": part["id"]}).json()["text"]

    assert client.get(url, params={"path": "不存在的章节"}).status_code == 404


def test_tender_text_reports_truncation(client, project_id, tmp_path):
    parsed = _setup(client, project_id, tmp_path)
    head = client.get(f"/api/v1/project/{project_id}/tender/text").json()
    assert head["total"] == head["length"] == len(parsed["full_text"])
    assert head["truncated"] == (len(parsed["full_text"]) > 5000)
    small = client.get(f"/api/v1/project/{project_id}/tender/text", params={"limit": 50}).json()
    assert small["truncated"] and len(small["text"]) == 50 and small["end"] == 50


def test_no_structure_for_old_projects(client, project_id):
    res = client.get(f"/api/v1/project/{project_id}/tender/outline").json()
    assert res == {"has_structure": False, "source": "", "sections": [], "keywords": [], "related": []}
    assert client.get(f"/api/v1/project/{project_id}/tender/section", params={"path": "sec_1"}).status_code == 404
