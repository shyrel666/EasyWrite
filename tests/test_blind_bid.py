"""暗标副本、身份跨 run/属性/关系检查、附件隔离与原项目保护。"""
import io
import os
import time
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from app.models.schemas import BlindBidConfig
from app.services.exporter.blind_bid import IMAGES, identifier_pattern, redact, scrub_docx
from app.services.project_store import find_node, project_store

MARK = '【已隐去】'


def _part(content, name='word/document.xml'):
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        return z.read(name).decode()


def _body_text(content):
    return ''.join(etree.fromstring(_part(content).encode()).itertext())


def test_blind_export_removes_identifiers_without_mutating_project(client, project_id):
    url = f'/api/v1/project/{project_id}'
    facts = {'company_name': '甲乙测试企业', 'legal_rep': '刘某某', 'credit_code': 'TEST91350012345',
             'custom_facts': {'办公地址': '测试街道99号', '收款账号': '12345678999999'}}
    client.put(url + '/facts', json=facts)
    content = '甲乙测试企业提供服务。刘某某在测试街道99号办公，账号12345678999999。\n\n[联系企业](https://example.com/甲乙测试企业)\n\n1. 实施\n2. 验收'
    client.put(url + '/outline', json={'outline': [{'id': 's', 'title': '服务方案', 'content': content}]})
    original = project_store.get(project_id)
    assert client.put(url + '/blind-bid', json={'enabled': True, 'font_size': 14, 'rules': '宋体四号'}).status_code == 200
    pre = client.get(url + '/export/preflight').json()['blind_bid']
    assert len(pre['hits']) == 4
    res = client.post(url + '/export', json={})
    assert res.status_code == 200
    with zipfile.ZipFile(io.BytesIO(res.content)) as z:
        all_xml = '\n'.join(z.read(n).decode() for n in z.namelist() if n.endswith(('.xml', '.rels')))
        for value in [facts['company_name'], facts['legal_rep'], facts['credit_code'], *facts['custom_facts'].values()]:
            assert value not in all_xml
        assert '投 标 单 位' not in all_xml
        root = etree.fromstring(z.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        assert set(root.xpath('//w:sz/@w:val', namespaces=ns)) == {'28'}
        assert not any(n.startswith('word/media/') for n in z.namelist())
    current = project_store.get(project_id)
    assert current.facts == original.facts and current.outline == original.outline
    if os.getenv('EASYWRITE_QA_DIR'):
        (Path(os.environ['EASYWRITE_QA_DIR']) / 'p1-8-blind.docx').write_bytes(res.content)


def test_blind_images_require_explicit_omission_and_separate_proofs(client, project_id):
    url = f'/api/v1/project/{project_id}'
    client.put(url + '/outline', json={'outline': [{'id': 's', 'title': '网络', 'content': '```mermaid\ngraph TD\nA[服务] --> B[数据]\n```'}]})
    client.put(url + '/blind-bid', json={'enabled': True})
    assert client.get(url + '/export/preflight').json()['blind_bid']['blocked']
    assert client.post(url + '/export', json={}).status_code == 400
    client.put(url + '/blind-bid', json={'enabled': True, 'omit_images': True})
    res = client.post(url + '/export', json={})
    assert res.status_code == 200
    with zipfile.ZipFile(io.BytesIO(res.content)) as z:
        assert not any(n.startswith('word/media/') for n in z.namelist())
    plan = [{'section_id': 's', 'asset_key': 'qualifications:q', 'attachment_id': 'a', 'sha256': '0' * 64}]
    assert client.post(url + '/export', json={'appendix_plan': plan}).status_code == 400
    assert client.post(url + '/export', json={'book_kind': 'qualification'}).status_code == 400


def test_split_runs_metadata_headers_and_template_default(tmp_path, client, project_id):
    path = tmp_path / 'source.docx'
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('甲乙')
    p.add_run('企业').bold = True
    doc.core_properties.author = '甲乙企业'
    doc.sections[0].header.paragraphs[0].text = '甲乙企业'
    doc.sections[0].footer.paragraphs[0].text = '甲乙企业'
    doc.part.relate_to('https://example.com/%E7%94%B2%E4%B9%99',
                       'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    doc.save(path)
    scrub_docx(path, BlindBidConfig(enabled=True), ['甲乙企业'])
    with zipfile.ZipFile(path) as z:
        text = '\n'.join(z.read(n).decode() for n in z.namelist() if n.endswith('.xml'))
        assert '甲乙企业' not in text
        rels = z.read('word/_rels/document.xml.rels').decode()
        assert 'about:blank' in rels and '%E7%94%B2' not in rels
    client.post('/api/v1/templates/create', json={'id': 'blindtest', 'name': '暗标格式', 'blind_bid': {'enabled': True}})
    res = client.get(f'/api/v1/project/{project_id}/export/preflight?template_id=blindtest')
    assert res.json()['blind_bid']['config']['enabled']


def test_structure_attributes_untouched_and_short_numbers_ignored(client, project_id):
    """自定义事实里的短数字不是标识；用户显式补充时只替换正文，缩进、行距等数值属性保持原值（否则 Word 判定损坏）。"""
    url = f'/api/v1/project/{project_id}'
    client.put(url + '/facts', json={'company_name': '甲乙测试企业', 'custom_facts': {'公司员工数': '360', '公司简称': '甲乙'}})
    content = '甲乙测试企业共360人，由甲乙负责实施。\n\n1. 实施\n2. 验收'
    client.put(url + '/outline', json={'outline': [{'id': 's', 'title': '服务方案', 'content': content}]})
    client.put(url + '/blind-bid', json={'enabled': True})
    pre = client.get(url + '/export/preflight').json()['blind_bid']
    assert [h['identifier'] for h in pre['hits']] == ['甲乙测试企业', '甲乙']
    body = _body_text(client.post(url + '/export', json={}).content)
    assert '360人' in body and '甲乙' not in body

    client.put(url + '/blind-bid', json={'enabled': True, 'identifiers': ['360']})
    res = client.post(url + '/export', json={})
    assert res.status_code == 200 and '360' not in _body_text(res.content)
    with zipfile.ZipFile(io.BytesIO(res.content)) as z:
        for name in z.namelist():
            if name.endswith(('.xml', '.rels')):
                root = etree.fromstring(z.read(name))
                assert not [v for e in root.iter(etree.Element) for v in e.attrib.values() if MARK in v], name
        numbering = etree.fromstring(z.read('word/numbering.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        assert '360' in numbering.xpath('//w:ind/@w:hanging', namespaces=ns) + numbering.xpath('//w:ind/@w:left', namespaces=ns)


def test_overlapping_identifiers_fully_redacted():
    """两个标识首尾交叠时整段隐去，不因替换先后残留半截。"""
    pattern = identifier_pattern(['甲乙', '乙测试企业有限'])
    assert redact('甲乙测试企业有限公司', pattern) == MARK + '公司'
    assert redact('甲 乙\n测试企业有限', pattern) == MARK
    assert redact('无标识正文', identifier_pattern([])) == '无标识正文'


def test_project_setting_overrides_template_default(client, project_id):
    """模板默认开启暗标时项目可显式关闭；enabled 为 null 时恢复沿用模板。"""
    url = f'/api/v1/project/{project_id}'
    client.put(url + '/facts', json={'company_name': '甲乙测试企业'})
    client.put(url + '/outline', json={'outline': [{'id': 's', 'title': '服务', 'content': '甲乙测试企业提供服务。'}]})
    client.post('/api/v1/templates/create', json={'id': 'blindtpl', 'name': '暗标格式', 'blind_bid': {'enabled': True}})
    preflight = url + '/export/preflight?template_id=blindtpl'
    assert client.get(preflight).json()['blind_bid']['config']['enabled'] is True

    assert client.put(url + '/blind-bid', json={'enabled': False}).status_code == 200
    assert client.get(preflight).json()['blind_bid']['config']['enabled'] is False
    xml = _part(client.post(url + '/export', json={'template_id': 'blindtpl'}).content)
    assert '投 标 单 位' in xml and '甲乙测试企业' in xml

    client.put(url + '/blind-bid', json={'enabled': None})
    assert client.get(preflight).json()['blind_bid']['config']['enabled'] is True
    xml = _part(client.post(url + '/export', json={'template_id': 'blindtpl'}).content)
    assert '投 标 单 位' not in xml and '甲乙测试企业' not in xml


def test_personnel_from_section_refs_are_redacted(client, project_id):
    """自动匹配写入章节的人员只记录在 last_refs，同样脱敏；资料已删除时用记录中的姓名。"""
    url = f'/api/v1/project/{project_id}'
    client.post('/api/v1/assets/personnel', json={'id': 'p_blind', 'name': '孙某某', 'role': '项目经理', 'status': 'confirmed'})
    client.put(url + '/outline', json={'outline': [{'id': 's', 'title': '项目团队', 'content': '拟派项目经理孙某某，技术负责人王某某。'}]})

    def add_refs(p):
        find_node(p.outline, 's').last_refs = [
            {'ref_type': 'asset', 'kind': 'personnel', 'asset_id': 'p_blind', 'name': '孙某某', 'source': 'matched'},
            {'ref_type': 'asset', 'kind': 'personnel', 'asset_id': 'p_deleted', 'name': '王某某', 'source': 'matched'},
        ]

    project_store.update(project_id, add_refs)
    client.put(url + '/blind-bid', json={'enabled': True})
    hits = {h['identifier'] for h in client.get(url + '/export/preflight').json()['blind_bid']['hits']}
    assert {'孙某某', '王某某'} <= hits
    body = _body_text(client.post(url + '/export', json={}).content)
    assert '孙某某' not in body and '王某某' not in body and '拟派项目经理' in body


def test_tilde_and_long_mermaid_fences_count_as_images(client, project_id):
    """与导出器一致：~~~ 围栏、四个反引号及未闭合围栏都算图代码，排除后可正常导出。"""
    assert IMAGES.sub('', '前\n````mermaid\ngraph TD\n````\n后') == '前\n\n后'
    assert IMAGES.sub('', '前\n```mermaid\ngraph TD\nA-->B') == '前\n'
    assert IMAGES.sub('', '```python\nprint(1)\n```') == '```python\nprint(1)\n```'
    assert IMAGES.sub('', '列表：\n   > ~~~mermaid\n   > graph TD\n   > ~~~\n后') == '列表：\n\n后'
    inline = '可用 ```mermaid``` 语法描述流程。\n后文保留'
    assert IMAGES.sub('', inline) == inline
    url = f'/api/v1/project/{project_id}'
    content = '说明\n\n~~~mermaid\ngraph TD\nA[服务] --> B[数据]\n~~~\n\n结论'
    client.put(url + '/outline', json={'outline': [{'id': 's', 'title': '网络', 'content': content}]})
    client.put(url + '/blind-bid', json={'enabled': True})
    pre = client.get(url + '/export/preflight').json()['blind_bid']
    assert pre['images'] and pre['blocked']
    client.put(url + '/blind-bid', json={'enabled': True, 'omit_images': True})
    res = client.post(url + '/export', json={})
    assert res.status_code == 200
    with zipfile.ZipFile(io.BytesIO(res.content)) as z:
        assert not any(n.startswith('word/media/') for n in z.namelist())
    body = _body_text(res.content)
    assert '说明' in body and '结论' in body and 'graph TD' not in body


def test_template_blind_rules_validated(client):
    """模板里的暗标规则按同一模型校验：非法值返回 422，null 视为未设置。"""
    for bad in [{'font_name': '微软雅黑'}, {'font_size': 99}, 'abc']:
        assert client.post('/api/v1/templates/create', json={'id': 'bad', 'name': '坏', 'blind_bid': bad}).status_code == 422
    res = client.post('/api/v1/templates/create', json={'id': 'nullbb', 'name': '空', 'blind_bid': None})
    assert res.status_code == 200 and res.json()['template']['blind_bid']['enabled'] is None


def test_saving_blind_settings_keeps_updated_at(client, project_id, monkeypatch):
    """暗标设置是导出格式，不刷新 updated_at，红线核查不会因此被判为过期；改正文仍会刷新。"""
    url = f'/api/v1/project/{project_id}'
    before = project_store.get(project_id).updated_at
    monkeypatch.setattr('app.services.project_store.now_str', lambda: '2999-01-01 00:00:00')
    assert client.put(url + '/blind-bid', json={'enabled': True}).status_code == 200
    assert project_store.get(project_id).updated_at == before
    client.put(url + '/facts', json={'company_name': '甲乙测试企业'})
    assert project_store.get(project_id).updated_at == '2999-01-01 00:00:00'


def test_redaction_scales_with_many_identifiers(tmp_path):
    """标识合成单个模式且只扫文字节点：200 个补充标识的长文档仍在数秒内完成（逐标识逐属性替换需数十秒）。"""
    path = tmp_path / 'big.docx'
    doc = Document()
    for i in range(3000):
        p = doc.add_paragraph('系统采用分层架构，接入层负责协议适配，网关统一鉴权；')
        p.add_run(f'服务层按领域拆分第{i}组，数据层采用主备复制。').bold = True
    doc.add_paragraph('联系甲乙企业')
    doc.save(path)
    identifiers = [f'别名标识{i:03d}' for i in range(200)] + ['甲乙企业']
    start = time.perf_counter()
    scrub_docx(path, BlindBidConfig(enabled=True), identifiers)
    elapsed = time.perf_counter() - start
    assert '甲乙企业' not in _part(path.read_bytes())
    assert elapsed < 5, elapsed  # 本机实测约 0.3 秒，旧实现约 27 秒
