"""同名/同项目并发导出必须隔离；下载完成清理文件，写入异常不留半成品。"""
import io
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Barrier
from urllib.parse import unquote

import pytest
from docx import Document
from docx.document import Document as DocumentClass

from app.core.config import settings
from app.models.schemas import DeviationItem, OutlineNode
from app.services.exporter.docx_generator import docx_exporter


def document_text(path_or_buffer):
    doc = Document(path_or_buffer)
    return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])


@pytest.mark.parametrize("deviation", [False, True])
@pytest.mark.parametrize("names", [("同名项目", "同名项目"), ("项目/A", "项目:A")])
def test_concurrent_export_paths_and_contents_are_isolated(deviation, names):
    barrier = Barrier(2)

    def export(i):
        text = f"请求{i}独有的正文"
        barrier.wait(timeout=5)
        if deviation:
            path = docx_exporter.export_deviation_table(names[i], [DeviationItem(index=1, clause_title=text)])
        else:
            path = docx_exporter.export_project_to_docx(names[i], "采购人", [OutlineNode(id="s1", title="方案", content=text)])
        return path

    with ThreadPoolExecutor(max_workers=2) as executor:
        paths = list(executor.map(export, range(2)))
    try:
        assert paths[0] != paths[1]
        for i, path in enumerate(paths):
            assert path.parent.resolve() == settings.OUTPUT_DIR.resolve()
            text = document_text(path)
            assert f"请求{i}独有的正文" in text and f"请求{1-i}独有的正文" not in text
    finally:
        for path in paths:
            path.unlink(missing_ok=True)


def test_explicit_filename_is_also_unique_and_stays_in_export_directory():
    paths = [docx_exporter.export_project_to_docx("项目", "采购人", [], output_filename="../../same.docx") for _ in range(2)]
    try:
        assert paths[0] != paths[1]
        assert all(p.resolve().parent == settings.OUTPUT_DIR.resolve() for p in paths)
    finally:
        for path in paths:
            path.unlink(missing_ok=True)


@pytest.mark.parametrize("deviation", [False, True])
def test_response_cleans_own_output_and_keeps_friendly_download_name(client, project_id, monkeypatch, deviation):
    method = "export_deviation_table" if deviation else "export_project_to_docx"
    real = getattr(docx_exporter, method)
    paths = []

    def capture(*args, **kwargs):
        path = real(*args, **kwargs)
        paths.append(path)
        assert path.exists()
        return path

    monkeypatch.setattr(docx_exporter, method, capture)
    client.put(f"/api/v1/project/{project_id}/outline", json={"outline": [{"id": "s", "title": "方案", "content": "独有内容"}]})
    client.put(f"/api/v1/project/{project_id}/deviation", json=[{"index": 1, "clause_title": "独有内容"}])
    sentinel = settings.OUTPUT_DIR / "existing-user-file.docx"
    sentinel.write_bytes(b"existing-user-file")
    suffix = "/deviation/export" if deviation else "/export"
    response = client.get(f"/api/v1/project/{project_id}{suffix}")
    assert response.status_code == 200 and "独有内容" in document_text(io.BytesIO(response.content))
    expected = docx_exporter.download_name(client.get(f"/api/v1/project/{project_id}").json()["name"],
                                            "技术偏离表" if deviation else "技术标书")
    assert expected in unquote(response.headers["content-disposition"])
    assert paths and all(not p.exists() for p in paths)
    assert sentinel.read_bytes() == b"existing-user-file"


def test_save_failure_removes_only_its_partial_file(monkeypatch):
    before = set(settings.OUTPUT_DIR.iterdir())

    def fail(_doc, filename):
        Path(filename).write_bytes(b"partial")
        raise OSError("模拟磁盘写入失败")

    monkeypatch.setattr(DocumentClass, "save", fail)
    with pytest.raises(OSError, match="模拟磁盘写入失败"):
        docx_exporter.export_project_to_docx("项目", "采购人", [])
    assert set(settings.OUTPUT_DIR.iterdir()) == before


@pytest.mark.parametrize("same_project", [False, True])
def test_overlapping_downloads_finish_before_their_files_are_cleaned(client, project_id, monkeypatch, same_project):
    other = project_id if same_project else client.post("/api/v1/project/create", json={"name": "并发项目"}).json()["id"]
    pids = [project_id, other]
    for i, pid in enumerate(dict.fromkeys(pids)):
        client.put(f"/api/v1/project/{pid}/outline", json={"outline": [{"id": "s", "title": "方案", "content": f"正文版本{i}"}]})
    real = docx_exporter.export_project_to_docx
    barrier = Barrier(2)
    paths = []

    def capture(*args, **kwargs):
        # 即便项目名相同且两个响应均在生成完毕后才读取文件，也必须各自正确。
        kwargs["project_name"] = "同名并发项目"
        path = real(*args, **kwargs)
        paths.append(path)
        barrier.wait(timeout=10)
        return path

    monkeypatch.setattr(docx_exporter, "export_project_to_docx", capture)
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda pid: client.get(f"/api/v1/project/{pid}/export"), pids))
        assert len(set(paths)) == 2
        for i, response in enumerate(responses):
            assert response.status_code == 200
            expected = 0 if same_project else i
            assert f"正文版本{expected}" in document_text(io.BytesIO(response.content))
        assert all(not path.exists() for path in paths)
    finally:
        if other != project_id:
            client.delete(f"/api/v1/project/{other}")
