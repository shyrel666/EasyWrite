"""证明原件的页序、资源边界、关联链路和 Word 中的保真插入。"""
import io
import os
from pathlib import Path
from zipfile import ZipFile

import fitz
import pytest
from docx import Document
from PIL import Image

from app.models.schemas import EvidenceAppendixEntry, OutlineNode, TenderAnalysis18
from app.services.assets.asset_manager import asset_manager
from app.services.exporter.docx_generator import docx_exporter
from app.services.exporter.evidence_appendix import AppendixError, appendix_options, prepare_appendices
from app.services.exporter.scoring_index import scoring_index_rows
from app.services.project_store import project_store


def pdf():
    with fitz.open() as doc:
        for i in range(2):
            page = doc.new_page(width=700, height=450)
            page.draw_rect(fitz.Rect(12, 12, 688, 438), color=(0, 0, 0), width=3)
            page.insert_text((40, 80), f'SYNTHETIC PROOF PAGE {i + 1}', fontsize=28)
            page.draw_circle(fitz.Point(610, 350), 40, color=(1, 0, 0), width=4)
        return doc.tobytes()


@pytest.fixture()
def setup(client, project_id):
    created = client.post('/api/v1/assets/qualifications', json={"id": "qa-proof", "name": "合成证书", "cert_no": "TEST-001", "issue_org": "合成测试机构", "expiry_date": "长期", "status": "confirmed", "holder": "甲公司"})
    assert created.status_code == 200, created.text
    response = client.post('/api/v1/assets/qualifications/qa-proof/attachments', files={"file": ("证明.pdf", pdf(), "application/pdf")})
    assert response.status_code == 200, response.text
    att = response.json()
    p = project_store.get(project_id)
    p.facts.company_name = "甲公司"
    p.tender_analysis = TenderAnalysis18(scoring_items=[{"id": "i1", "name": "企业资质", "response_type": "evidence"}, {"id": "i2", "name": "团队资质", "response_type": "evidence"}])
    p.evidence_links = {"i1": ["qualifications:qa-proof"], "i2": ["qualifications:qa-proof"]}
    p.outline = [OutlineNode(id="s", title="第一章 资格证明", content="证明资料见附件。", scoring_item_ids=["i1", "i2"])]
    return p, att


def entry(p, **updates):
    option = next(o for o in appendix_options(p) if o.get('available'))
    return EvidenceAppendixEntry(**{**option, **updates})


def test_pdf_page_order_shared_links_and_original_borders(setup):
    p, _ = setup
    plan = entry(p, pages=[2, 1])
    prepared = prepare_appendices(p, [plan, plan])
    assert len(prepared) == 1 and [n for n, _ in prepared[0].images] == [2, 1]
    assert prepared[0].item_ids == ["i1", "i2"]
    with Image.open(io.BytesIO(prepared[0].images[0][1])) as image:
        assert abs(image.width / image.height - 700 / 450) < .01
    path = docx_exporter.export_project_to_docx("证明附件回归", "采购人", p.outline, facts=p.facts,
                                               scoring_index=scoring_index_rows(p), evidence_appendices=prepared)
    try:
        doc = Document(path)
        assert len(doc.inline_shapes) == 2
        titles = [x.text for x in doc.paragraphs if '原件第' in x.text]
        assert '第 2 页' in titles[0] and '第 1 页' in titles[1]
        with ZipFile(path) as z:
            xml = z.read('word/document.xml').decode()
            assert prepared[0].bookmark in xml and xml.count('PAGEREF') >= 6
            assert '未编入' not in '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        if os.getenv('EASYWRITE_QA_DIR'):
            (Path(os.environ['EASYWRITE_QA_DIR']) / 'p1-5.docx').write_bytes(path.read_bytes())
    finally:
        path.unlink(missing_ok=True)


@pytest.mark.parametrize('pages', [[0], [3], [1, 1]])
def test_missing_or_duplicate_pages_rejected(setup, pages):
    p, _ = setup
    with pytest.raises(AppendixError, match='页序'):
        prepare_appendices(p, [entry(p, pages=pages)])


def test_subject_and_attachment_fingerprint_changes_rejected(setup):
    p, att = setup
    plan = entry(p)
    p.facts.company_name = '乙公司'
    assert not appendix_options(p)[0]['available']
    with pytest.raises(AppendixError, match='主体'):
        prepare_appendices(p, [plan])
    p.facts.company_name = '甲公司'
    path, _ = asset_manager.attachment_file('qualifications', 'qa-proof', att['id'])
    path.write_bytes(path.read_bytes() + b' changed')
    with pytest.raises(AppendixError, match='指纹'):
        prepare_appendices(p, [plan])


def test_registered_path_must_stay_under_asset_storage(setup, monkeypatch, tmp_path):
    p, att = setup
    plan = entry(p)
    _, meta = asset_manager.attachment_file('qualifications', 'qa-proof', att['id'])
    outside = tmp_path / 'outside.pdf'
    outside.write_bytes(pdf())
    monkeypatch.setattr(asset_manager, 'attachment_file', lambda *args: (outside, meta))
    with pytest.raises(AppendixError, match='路径越界'):
        prepare_appendices(p, [plan])


def test_broken_attachment_and_section_mapping(setup, client):
    p, _ = setup
    with pytest.raises(AppendixError, match='章节'):
        prepare_appendices(p, [entry(p, placement='section', section_id='missing')])
    client.post('/api/v1/assets/qualifications/qa-proof/attachments', files={"file": ("bad.pdf", b'%PDF-invalid', 'application/pdf')})
    options = appendix_options(p)
    assert not options[-1]['available'] and any('损坏' in s for s in options[-1]['problems'])


def test_pixel_and_total_page_limits(setup, monkeypatch):
    from app.services.exporter import evidence_appendix as ea
    p, _ = setup
    monkeypatch.setattr(ea, 'MAX_PAGES', 1)
    with pytest.raises(AppendixError, match='100 页'):
        prepare_appendices(p, [entry(p)])


def test_image_pixels_and_partial_page_warning(setup, monkeypatch, client):
    from app.services.exporter import evidence_appendix as ea
    p, _ = setup
    partial = prepare_appendices(p, [entry(p, pages=[2])])
    assert '未编入原件页：1' in partial[0].warnings
    image = Image.new('RGB', (800, 400), 'white')
    buf = io.BytesIO()
    image.save(buf, format='PNG')
    result = client.post('/api/v1/assets/qualifications/qa-proof/attachments', files={"file": ("横向扫描.png", buf.getvalue(), 'image/png')})
    assert result.status_code == 200
    option = appendix_options(p)[-1]
    plan = EvidenceAppendixEntry(**option)
    prepared = prepare_appendices(p, [plan])
    with Image.open(io.BytesIO(prepared[0].images[0][1])) as output:
        assert output.size == (800, 400)
    monkeypatch.setattr(ea, 'MAX_PIXELS', 100)
    with pytest.raises(AppendixError, match='像素'):
        prepare_appendices(p, [plan])
