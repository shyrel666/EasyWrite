import os
import sys
from pathlib import Path
import docx
from docx import Document

# 将 backend 路径加入 sys.path
BASE_DIR = Path(__file__).resolve().parent.parent / "backend"
sys.path.insert(0, str(BASE_DIR))

from app.services.parser.word_parser import WordDocumentParser
from app.services.rag.vector_store import HierarchicalVectorStore
from app.services.generator.bid_generator import BidGenerator
from app.services.exporter.docx_generator import DocxBidExporter
from app.models.schemas import OutlineNode

def create_sample_docx(file_path: Path):
    """动态生成一份逼真的信息化历史中标技术方案 Word 样本"""
    doc = Document()
    
    # 标题 1
    p1 = doc.add_paragraph("第一章 某市政务大数据平台建设方案")
    p1.style = 'Heading 1'
    doc.add_paragraph("本项目旨在打造全市统一政务数据中台，消除信息孤岛，实现跨部门业务协同。")

    # 标题 2
    p1_1 = doc.add_paragraph("1.1 建设背景与痛点分析")
    p1_1.style = 'Heading 2'
    doc.add_paragraph("当前各委办局业务系统分散，数据标准不一，缺乏统一的数据资产目录与共享交换通道。")

    # 标题 1
    p2 = doc.add_paragraph("第二章 系统总体技术架构设计")
    p2.style = 'Heading 1'
    doc.add_paragraph("系统遵循分层解耦、松耦合与高可用设计原则，采用微服务架构与云原生部署模式。")

    # 标题 2
    p2_1 = doc.add_paragraph("2.1 微服务架构与容器化部署")
    p2_1.style = 'Heading 2'
    doc.add_paragraph("服务层采用 Spring Cloud 与 Kubernetes 容器集群，集成 Istio 服务网格，支持灰度发布与秒级自动弹性伸缩。")

    # 插入技术选型表格
    table = doc.add_table(rows=3, cols=3)
    data = [
        ["模块名称", "选型技术", "技术规格与特性"],
        ["API 网关", "Spring Cloud Gateway", "支持双向 SSL、动态路由、分布式限流"],
        ["服务治理", "Kubernetes + Istio", "容器编排、全链路追踪与无感知热升级"]
    ]
    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val

    # 标题 2
    p2_2 = doc.add_paragraph("2.2 数据库容灾备份与高可用方案")
    p2_2.style = 'Heading 2'
    doc.add_paragraph("数据库采用主从复制 + 半同步机制，支持同城双活与异地灾备，RPO=0，RTO<30秒，满足等保三级核心数据备份要求。")

    # 标题 1
    p3 = doc.add_paragraph("第三章 售后服务保障体系")
    p3.style = 'Heading 1'
    doc.add_paragraph("我司设立 7×24 小时现场与远程响应中心，重大故障 15 分钟内响应，1 小时内到达现场。")

    file_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(file_path))
    print(f"[1] 测试样本 Word 已生成: {file_path}")

def test_full_pipeline():
    test_dir = Path(__file__).resolve().parent / "test_data"
    test_dir.mkdir(parents=True, exist_ok=True)
    sample_file = test_dir / "sample_bid.docx"
    
    # 步骤 1：生成样本 Word
    create_sample_docx(sample_file)
    assert sample_file.exists(), "样本文件生成失败"

    # 步骤 2：测试大纲层级解析器
    parser = WordDocumentParser()
    parsed_res = parser.parse_docx(sample_file)
    print(f"[2] 文档解析成功，总章节数: {parsed_res['total_sections']}")
    assert len(parsed_res["sections"]) >= 3, "一级大纲数量不符合预期"
    
    # 校验面包屑切片
    chunks = parser.flatten_sections_to_chunks(parsed_res, doc_name="sample_bid.docx")
    print(f"    扁平化切片总数: {len(chunks)}")
    assert len(chunks) >= 4, "切片数量不足"
    
    # 验证是否有包含微服务和大纲路径的切片
    ms_chunk = next((c for c in chunks if "微服务" in c["section_title"]), None)
    assert ms_chunk is not None, "未找到微服务对应切片"
    print(f"    切片大纲面包屑路径验证: {ms_chunk['breadcrumb']}")
    assert "第二章 系统总体技术架构设计" in ms_chunk["breadcrumb"]

    # 步骤 3：测试知识库入库与混合检索
    kb_dir = test_dir / "kb_store"
    kb_dir.mkdir(parents=True, exist_ok=True)
    kb = HierarchicalVectorStore(storage_dir=kb_dir)
    kb.clear()
    added = kb.add_chunks(chunks)
    print(f"[3] 知识库入库成功，导入切片数: {added}")
    
    # 执行混合检索
    query = "请问系统如何进行微服务容器化与容器集群部署？"
    search_results = kb.search(query, top_k=2)
    print(f"    检索命中结果数: {len(search_results)}")
    assert len(search_results) > 0, "检索结果为空"
    top_hit = search_results[0]
    print(f"    最高相关性切片: [{top_hit['score']}] {top_hit['breadcrumb']}")
    assert "微服务" in top_hit["content"] or "微服务" in top_hit["section_title"]

    # 步骤 4：测试分章节智能草拟
    gen = BidGenerator()
    gen.kb = kb  # 挂载测试知识库
    print("[4] 开始模拟分章草拟...")
    draft = gen.draft_section(
        section_title="2.1 系统微服务治理方案",
        section_path="总体方案设计 > 系统微服务治理方案",
        requirements=["必须说明微服务组件选型", "阐明高可用与弹性伸缩机制"],
        custom_instruction="重点强调容器化与双活架构"
    )
    print(f"    生成草稿字数: {len(draft['generated_content'])}")
    print(f"    引用参考资产数: {len(draft['references'])}")
    assert len(draft["generated_content"]) > 50, "生成内容过短"

    # 步骤 5：测试高保真 Word 导出
    print("[5] 测试将大纲树高保真导出为标准标书 Word...")
    test_outline = [
        OutlineNode(
            id="sec_1",
            title="第一章 项目总体规划与微服务架构设计",
            level=1,
            path="第一章 项目总体规划与微服务架构设计",
            content=draft["generated_content"],
            children=[
                OutlineNode(
                    id="sec_1_1",
                    title="1.1 技术选型对照表",
                    level=2,
                    path="第一章 项目总体规划与微服务架构设计 > 1.1 技术选型对照表",
                    content=(
                        "以下为本次投标建议采用的核心技术选型清单：\n\n"
                        "| 技术类别 | 选型组件 | 响应能力说明 |\n"
                        "| --- | --- | --- |\n"
                        "| 分布式网关 | Spring Cloud Gateway | 完全满足毫秒级路由转发与鉴权 |\n"
                        "| 数据库中间件 | ShardingSphere | 支持分库分表与读写分离 |\n"
                        "| 容灾与备份 | 异地双活集群 | RPO=0，满足等保三级要求 |\n\n"
                        "**我方承诺**：所有组件均经过充分兼容性测试，提供 5 年免费技术升级与维护。"
                    )
                )
            ]
        )
    ]

    exporter = DocxBidExporter()
    out_file = exporter.export_project_to_docx(
        project_name="市级政务大数据中台建设项目",
        client_name="市政务服务和数字化建设管理局",
        outline=test_outline,
        output_filename="测试输出_市级政务大数据中台技术标书.docx"
    )
    print(f"[5] 标书导出成功: {out_file}")
    assert out_file.exists(), "导出标书文件未生成"
    assert out_file.stat().st_size > 10000, "导出的 Word 文件体积异常（过小）"

    print("\n[SUCCESS] Full pipeline test passed successfully!")

if __name__ == "__main__":
    if sys.platform == "win32":
        import io
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    test_full_pipeline()
