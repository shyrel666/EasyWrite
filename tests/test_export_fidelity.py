"""
Word 导出保真度：
- 正文里的 # 小标题排到所属章节下一级（真实标题样式，不与大纲章节平级）；整段加粗短句是不进目录的小标题
- 编号/项目符号列表为 Word 原生多级编号，每个列表各自从头编号，起始序号沿用
- Markdown 表格为原生表格：表头跨页重复、加粗生效（不残留 **）、短列不被均分撑宽
- A4 版面；中文正文单换行即分段
- 架构图优先用前端 Mermaid 渲染的图片，图题按全文编号；渲染不了的给代码插槽，绝不画无关的"缺省架构图"
- 接口：POST 导出带图片、偏离表单独导出
"""
import base64
import io

import docx as docxlib
from docx.oxml.ns import qn
from PIL import Image

from app.models.schemas import DeviationItem, OutlineNode
from app.services.exporter.diagram_renderer import diagram_renderer
from app.services.exporter.docx_generator import docx_exporter

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
FLOW = "graph TD\n    A[用户门户] --> B[API网关]\n    B --> C[业务微服务]"


def _png(width=400, height=200) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (200, 220, 255)).save(buf, format="PNG")
    return buf.getvalue()


def _export(outline, **kwargs):
    path = docx_exporter.export_project_to_docx("导出测试", "采购人", outline, **kwargs)
    return docxlib.Document(str(path))


def _body_after_toc(doc):
    paras = doc.paragraphs
    start = next(i for i, p in enumerate(paras) if p.style.name.startswith("Heading"))
    return paras[start:]


def _num_pr(p):
    num = p._p.find(f"{W_NS}pPr/{W_NS}numPr")
    if num is None:
        return None
    return int(num.find(f"{W_NS}numId").get(qn("w:val"))), int(num.find(f"{W_NS}ilvl").get(qn("w:val")))


def test_content_headings_nest_under_section():
    content = "# 总体设计\n\n正文一。\n\n## 部署方案\n\n正文二。\n\n**1.1.1 设计原则**\n\n正文三。"
    doc = _export([OutlineNode(id="a", title="第一章 方案", level=1, children=[
        OutlineNode(id="b", title="1.1 架构", level=2, content=content)])])
    styles = [(p.style.name, p.text) for p in _body_after_toc(doc)]
    assert ("Heading 2", "1.1 架构") in styles
    assert ("Heading 3", "总体设计") in styles  # 正文最浅的 # 标题 = 章节层级 + 1
    assert ("Heading 4", "部署方案") in styles
    bold_title = next(p for p in doc.paragraphs if p.text == "1.1.1 设计原则")
    assert bold_title.style.name == "Normal" and all(r.bold for r in bold_title.runs)
    assert bold_title.paragraph_format.first_line_indent == 0
    section = doc.sections[0]
    assert round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7


def test_native_numbering_and_softbreak_paragraphs():
    content = ("第一行正文\n第二行正文\n\n"
               "1. 先进性\n2. 可靠性\n   - 双活数据中心\n   - 自动切换\n3. 安全性\n\n"
               "过渡段落。\n\n"
               "3. 第三步\n4. 第四步\n\n"
               "- 项目经理\n- 实施工程师\n")
    doc = _export([OutlineNode(id="a", title="第一章", level=1, content=content)])
    paras = {p.text: p for p in doc.paragraphs}
    assert "第一行正文" in paras and "第二行正文" in paras  # 单换行分段
    assert "1. 先进性" not in paras and "先进性" in paras  # 编号由 Word 生成，不留在文字里

    first, nested, third = _num_pr(paras["先进性"]), _num_pr(paras["双活数据中心"]), _num_pr(paras["安全性"])
    assert first[1] == 0 and nested == (first[0], 1) and third == first
    second_list = _num_pr(paras["第三步"])
    assert second_list[0] != first[0]  # 另一个列表从头编号
    bullets = _num_pr(paras["项目经理"])
    assert bullets[0] not in (first[0], second_list[0])

    numbering = doc.part.numbering_part.element
    def abstract_of(num_id):
        num = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_id))
        aid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        return next(a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == aid)
    lvl = lambda a, i: a.findall(qn("w:lvl"))[i]  # noqa: E731
    first_abs = abstract_of(first[0])
    assert lvl(first_abs, 0).find(qn("w:numFmt")).get(qn("w:val")) == "decimal"
    assert lvl(first_abs, 1).find(qn("w:numFmt")).get(qn("w:val")) == "bullet"
    assert lvl(abstract_of(second_list[0]), 0).find(qn("w:start")).get(qn("w:val")) == "3"
    assert lvl(abstract_of(bullets[0]), 0).find(qn("w:numFmt")).get(qn("w:val")) == "bullet"


def test_markdown_table_is_native_and_formatted():
    content = ("| 序号 | 招标文件技术规格要求 | 投标响应状态 |\n| :---: | :--- | :---: |\n"
               "| 1 | 系统需支持国密SM4加密，并通过商用密码应用安全性评估 | **完全满足** |\n"
               "| 2 | 提供7×24小时运维响应 | **正偏离** |\n")
    doc = _export([OutlineNode(id="a", title="偏离表", level=1, content=content)])
    (table,) = doc.tables
    assert [c.text for c in table.rows[1].cells] == ["1", "系统需支持国密SM4加密，并通过商用密码应用安全性评估", "完全满足"]
    assert all("**" not in c.text for row in table.rows for c in row.cells)
    status = table.rows[1].cells[2].paragraphs[0]
    assert status.runs[0].bold and status.alignment == 1  # 居中
    assert table.rows[0]._tr.find(f"{W_NS}trPr/{W_NS}tblHeader") is not None  # 表头跨页重复
    widths = [c.width for c in table.rows[0].cells]
    assert widths[0] < widths[2] < widths[1]


def test_frontend_diagram_images_and_numbered_captions():
    png = _png()
    content = (f"```mermaid\n---\ntitle: 系统总体架构\n---\n{FLOW}\n```\n\n"
               "```mermaid\nsequenceDiagram\n  A->>B: 请求\n```\n")
    # 前端代码块缩进与后端不同也能匹配（逐行归一）
    doc = _export([OutlineNode(id="a", title="第一章 架构", level=1, content=content)],
                  diagrams={f"---\ntitle: 系统总体架构\n---\n{FLOW}".replace("    ", "  "): png})
    blobs = [r.target_part.blob for r in doc.part.rels.values() if "image" in r.reltype]
    assert blobs == [png]  # 只有前端图；时序图没有图片时不画替代图
    texts = [p.text for p in doc.paragraphs]
    assert "图 1  系统总体架构" in texts and "图 2  架构" in texts
    slot = doc.tables[0].cell(0, 0).text
    assert "架构图未能自动渲染" in slot and "sequenceDiagram" in slot


def test_fallback_renderer_never_invents_diagram():
    assert diagram_renderer._parse_mermaid("sequenceDiagram\n  A->>B: 请求") == ({}, [])
    assert diagram_renderer.render_to_image("mindmap\n  root((标书))") is None
    nodes, edges = diagram_renderer._parse_mermaid(FLOW + "\n    subgraph 核心\n    end")
    assert nodes == {"A": "用户门户", "B": "API网关", "C": "业务微服务"}  # 只引用 ID 不覆盖标签，指令行不是节点
    assert len(edges) == 2


def test_export_api_with_diagrams_and_deviation_table(client, project_id):
    outline = [{"id": "s1", "title": "第一章 架构", "level": 1, "content": f"```mermaid\n{FLOW}\n```"}]
    res = client.put(f"/api/v1/project/{project_id}/outline", json={"outline": outline})
    assert res.status_code == 200, res.text

    png = _png()
    data_url = "data:image/png;base64," + base64.b64encode(png).decode()
    res = client.post(f"/api/v1/project/{project_id}/export",
                      json={"template_id": "gov_red", "diagrams": [{"code": FLOW, "image": data_url}]})
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    doc = docxlib.Document(io.BytesIO(res.content))
    assert [r.target_part.blob for r in doc.part.rels.values() if "image" in r.reltype] == [png]

    bad = client.post(f"/api/v1/project/{project_id}/export",
                      json={"diagrams": [{"code": FLOW, "image": "data:image/png;base64,bm90IGEgcG5n"}]})
    assert bad.status_code == 400

    empty = client.get(f"/api/v1/project/{project_id}/deviation/export")
    assert empty.status_code == 400
    items = [
        DeviationItem(index=1, clause_title="※至少150人月", level="redline", response_status="完全满足",
                      response_detail="承诺投入150人月").model_dump(),
        DeviationItem(index=2, clause_title="支持国密SM4", level="normal", response_status="待生成").model_dump(),
    ]
    assert client.put(f"/api/v1/project/{project_id}/deviation", json=items).status_code == 200
    res = client.get(f"/api/v1/project/{project_id}/deviation/export")
    assert res.status_code == 200
    doc = docxlib.Document(io.BytesIO(res.content))
    (table,) = doc.tables
    assert len(table.rows) == 3
    assert [c.text for c in table.rows[1].cells] == ["1", "※至少150人月", "废标红线", "完全满足", "承诺投入150人月"]
    assert table.rows[1].cells[1].paragraphs[0].runs[0].bold  # 红线条款加粗
    assert "尚有 1 条未填写响应" in doc.paragraphs[-1].text
